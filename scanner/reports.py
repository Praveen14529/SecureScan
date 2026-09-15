import io
from django.utils import timezone
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from docx import Document
from docx.shared import Pt, RGBColor

LEVEL_HEX = {
    "low": "#1a7f37",
    "medium": "#9a6700",
    "high": "#bc4c00",
    "critical": "#cf222e",
}
LEVEL_COLORS = {key: colors.HexColor(hexcode) for key, hexcode in LEVEL_HEX.items()}
LEVEL_COLORS_RGB = {
    "low": RGBColor(0x1A, 0x7F, 0x37),
    "medium": RGBColor(0x9A, 0x67, 0x00),
    "high": RGBColor(0xBC, 0x4C, 0x00),
    "critical": RGBColor(0xCF, 0x22, 0x2E),
}
SEVERITY_LABELS = {"critical": "Critical", "high": "High", "medium": "Medium", "low": "Low", "info": "Info"}


def _scan_context(scan):
    level_key, level_label = scan.risk_level()
    findings = list(scan.findings.all())
    return {
        "scan": scan,
        "score": scan.risk_score(),
        "level_key": level_key,
        "level_label": level_label,
        "findings": findings,
        "breakdown": scan.category_breakdown(),
        "generated": timezone.now().strftime("%d %b %Y, %H:%M"),
    }


def render_scan_pdf(scan):
    ctx = _scan_context(scan)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleX", parent=styles["Title"], fontSize=18)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], spaceBefore=14)
    normal = styles["Normal"]

    story = [Paragraph("Vulnerability Scan Report", title_style), Spacer(1, 6)]

    meta_table = Table([
        ["Host", ctx["scan"].hostname or "-"],
        ["Operating System", ctx["scan"].os_summary or "-"],
        ["Scan Date", ctx["scan"].started_at.strftime("%d %b %Y, %H:%M")],
        ["Deep CVE Lookup", "Yes" if ctx["scan"].deep_cve_lookup else "No"],
        ["Generated", ctx["generated"]],
    ], colWidths=[4.5 * cm, 10.5 * cm])
    meta_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#555555")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#e0e0e0")),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 16))

    level_color = LEVEL_COLORS.get(ctx["level_key"], colors.grey)
    score_table = Table([
        ["Overall Risk Score", "Risk Level", "Total Findings"],
        [f"{ctx['score']} / 100", ctx["level_label"], str(len(ctx["findings"]))],
    ], colWidths=[5 * cm, 5 * cm, 5 * cm])
    score_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("FONTSIZE", (0, 1), (-1, 1), 14),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TEXTCOLOR", (1, 1), (1, 1), level_color),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dddddd")),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(score_table)
    story.append(Spacer(1, 16))

    story.append(Paragraph("Findings", h2))
    current_category = None
    for f in sorted(ctx["findings"], key=lambda x: (x.category, -x.severity_rank())):
        if f.category != current_category:
            current_category = f.category
            story.append(Paragraph(f.get_category_display(), styles["Heading3"]))
        sev_hex = LEVEL_HEX.get(f.severity, "#555555")
        story.append(Paragraph(
            f'<font color="{sev_hex}"><b>[{SEVERITY_LABELS.get(f.severity, f.severity)}]</b></font> {f.title}',
            normal,
        ))
        if f.description:
            story.append(Paragraph(f.description, ParagraphStyle("desc", parent=normal, fontSize=9,
                                                                    textColor=colors.HexColor("#444444"))))
        if f.mitigation:
            story.append(Paragraph(f"<b>Mitigation:</b> {f.mitigation}",
                                    ParagraphStyle("mit", parent=normal, fontSize=9,
                                                    textColor=colors.HexColor("#1a5c1a"))))
        story.append(Spacer(1, 8))

    doc.build(story)
    buffer.seek(0)
    return buffer


def render_scan_docx(scan):
    ctx = _scan_context(scan)
    document = Document()
    document.add_heading("Vulnerability Scan Report", level=0)

    meta = document.add_table(rows=0, cols=2)
    meta.style = "Light Grid Accent 1"
    for label, value in [
        ("Host", ctx["scan"].hostname or "-"),
        ("Operating System", ctx["scan"].os_summary or "-"),
        ("Scan Date", ctx["scan"].started_at.strftime("%d %b %Y, %H:%M")),
        ("Deep CVE Lookup", "Yes" if ctx["scan"].deep_cve_lookup else "No"),
        ("Generated", ctx["generated"]),
    ]:
        row = meta.add_row().cells
        row[0].text = label
        row[1].text = str(value)

    document.add_paragraph()
    p = document.add_paragraph()
    run = p.add_run(
        f"Overall Risk Score: {ctx['score']} / 100   |   Risk Level: {ctx['level_label']}   |   "
        f"Total Findings: {len(ctx['findings'])}"
    )
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = LEVEL_COLORS_RGB.get(ctx["level_key"], RGBColor(0, 0, 0))

    document.add_heading("Findings", level=1)
    current_category = None
    for f in sorted(ctx["findings"], key=lambda x: (x.category, -x.severity_rank())):
        if f.category != current_category:
            current_category = f.category
            document.add_heading(f.get_category_display(), level=2)
        p = document.add_paragraph()
        run = p.add_run(f"[{SEVERITY_LABELS.get(f.severity, f.severity)}] ")
        run.bold = True
        run.font.color.rgb = LEVEL_COLORS_RGB.get(f.severity, RGBColor(0, 0, 0))
        p.add_run(f.title)
        if f.description:
            document.add_paragraph(f.description)
        if f.mitigation:
            mp = document.add_paragraph()
            mp.add_run("Mitigation: ").bold = True
            mp.add_run(f.mitigation)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
