import io
from django.utils import timezone
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
)

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

LEVEL_COLORS = {
    "low": colors.HexColor("#1a7f37"),
    "medium": colors.HexColor("#9a6700"),
    "high": colors.HexColor("#bc4c00"),
    "critical": colors.HexColor("#cf222e"),
}
LEVEL_COLORS_RGB = {
    "low": RGBColor(0x1A, 0x7F, 0x37),
    "medium": RGBColor(0x9A, 0x67, 0x00),
    "high": RGBColor(0xBC, 0x4C, 0x00),
    "critical": RGBColor(0xCF, 0x22, 0x2E),
}


def _assessment_context(assessment):
    level_key, level_label = assessment.risk_level()
    return {
        "assessment": assessment,
        "asset": assessment.asset,
        "score": assessment.risk_score(),
        "compliance": assessment.compliance_percentage(),
        "level_key": level_key or "medium",
        "level_label": level_label,
        "breakdown": assessment.category_breakdown(),
        "responses": assessment.responses.select_related("question", "question__category").all(),
        "generated": timezone.now().strftime("%d %b %Y, %H:%M"),
    }


# ---------------------------------------------------------------- PDF -----
def render_assessment_pdf(assessment):
    ctx = _assessment_context(assessment)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleX", parent=styles["Title"], fontSize=18)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], spaceBefore=14)
    normal = styles["Normal"]

    story = []
    story.append(Paragraph("Cyber Risk Assessment Report", title_style))
    story.append(Paragraph(ctx["assessment"].title, styles["Heading3"]))
    story.append(Spacer(1, 6))

    meta_table = Table([
        ["Asset", ctx["asset"].name],
        ["Asset Type", ctx["asset"].get_asset_type_display()],
        ["Criticality", ctx["asset"].get_criticality_display()],
        ["Owner", ctx["asset"].owner or "-"],
        ["Assessor", ctx["assessment"].assessor or "-"],
        ["Status", ctx["assessment"].get_status_display()],
        ["Generated", ctx["generated"]],
    ], colWidths=[4 * cm, 11 * cm])
    meta_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#555555")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#e0e0e0")),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 16))

    # Score summary
    score = ctx["score"]
    level_color = LEVEL_COLORS.get(ctx["level_key"], colors.grey)
    score_table = Table([
        ["Overall Risk Score", "Compliance", "Risk Level"],
        [f"{score if score is not None else 'N/A'} / 100",
         f"{ctx['compliance'] if ctx['compliance'] is not None else 'N/A'}%",
         ctx["level_label"]],
    ], colWidths=[5 * cm, 5 * cm, 5 * cm])
    score_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("FONTSIZE", (0, 1), (-1, 1), 14),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TEXTCOLOR", (2, 1), (2, 1), level_color),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dddddd")),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(score_table)
    story.append(Spacer(1, 16))

    # Category breakdown
    if ctx["breakdown"]:
        story.append(Paragraph("Risk by Category", h2))
        data = [["Category", "Compliance %", "Risk %"]]
        for row in ctx["breakdown"]:
            data.append([row["category"], f'{row["compliance_pct"]}%', f'{row["risk_pct"]}%'])
        cat_table = Table(data, colWidths=[8 * cm, 4 * cm, 4 * cm])
        cat_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2b2b2b")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dddddd")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f7f7")]),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(cat_table)
        story.append(Spacer(1, 16))

    # Detailed responses
    story.append(Paragraph("Checklist Detail", h2))
    current_category = None
    for r in ctx["responses"]:
        if r.question.category_id != current_category:
            current_category = r.question.category_id
            story.append(Paragraph(r.question.category.name, styles["Heading4"]))
        story.append(Paragraph(f"<b>{r.question.text}</b>", normal))
        story.append(Paragraph(
            f"Answer: {r.get_answer_display()} (weight {r.question.weight})"
            + (f" &mdash; Notes: {r.evidence_notes}" if r.evidence_notes else ""),
            ParagraphStyle("resp", parent=normal, textColor=colors.HexColor("#444444"), fontSize=9),
        ))
        story.append(Spacer(1, 6))

    if ctx["assessment"].notes:
        story.append(PageBreak())
        story.append(Paragraph("Assessor Notes", h2))
        story.append(Paragraph(ctx["assessment"].notes, normal))

    doc.build(story)
    buffer.seek(0)
    return buffer


# --------------------------------------------------------------- DOCX -----
def render_assessment_docx(assessment):
    ctx = _assessment_context(assessment)
    document = Document()

    title = document.add_heading("Cyber Risk Assessment Report", level=0)
    document.add_heading(ctx["assessment"].title, level=2)

    meta = document.add_table(rows=0, cols=2)
    meta.style = "Light Grid Accent 1"
    for label, value in [
        ("Asset", ctx["asset"].name),
        ("Asset Type", ctx["asset"].get_asset_type_display()),
        ("Criticality", ctx["asset"].get_criticality_display()),
        ("Owner", ctx["asset"].owner or "-"),
        ("Assessor", ctx["assessment"].assessor or "-"),
        ("Status", ctx["assessment"].get_status_display()),
        ("Generated", ctx["generated"]),
    ]:
        row = meta.add_row().cells
        row[0].text = label
        row[1].text = str(value)

    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(
        f"Overall Risk Score: {ctx['score'] if ctx['score'] is not None else 'N/A'} / 100   |   "
        f"Compliance: {ctx['compliance'] if ctx['compliance'] is not None else 'N/A'}%   |   "
        f"Risk Level: {ctx['level_label']}"
    )
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = LEVEL_COLORS_RGB.get(ctx["level_key"], RGBColor(0, 0, 0))

    if ctx["breakdown"]:
        document.add_heading("Risk by Category", level=1)
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Category", "Compliance %", "Risk %"
        for row in ctx["breakdown"]:
            cells = table.add_row().cells
            cells[0].text = row["category"]
            cells[1].text = f'{row["compliance_pct"]}%'
            cells[2].text = f'{row["risk_pct"]}%'

    document.add_heading("Checklist Detail", level=1)
    current_category = None
    for r in ctx["responses"]:
        if r.question.category_id != current_category:
            current_category = r.question.category_id
            document.add_heading(r.question.category.name, level=2)
        p = document.add_paragraph()
        p.add_run(r.question.text).bold = True
        note = f" — Notes: {r.evidence_notes}" if r.evidence_notes else ""
        document.add_paragraph(f"Answer: {r.get_answer_display()} (weight {r.question.weight}){note}")

    if ctx["assessment"].notes:
        document.add_heading("Assessor Notes", level=1)
        document.add_paragraph(ctx["assessment"].notes)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
